#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import gzip
import os
import re
import sys
import traceback

import numpy as np
import pandas as pd


# ========================= Helpers =========================

TARGET_GENES = [
    "BRCA1", "BRCA2", "RAD51C", "RAD51D", "PMS2", "PALB2",
    "PTEN", "TP53", "CDH1", "MLH1", "MSH2", "MSH6", "EPCAM"
]
TARGET_GENE_SET = set(TARGET_GENES)


DESIRED_ORDER = [
    "pré-notes_variant", "notes_variant", "notes_échantillon",
    "Validation", "Sample_ID", "Gene_name", "ID", "SV_type", "SV_length",
    "REF", "ALT", "QUAL", "SUPPORT", "COVERAGE", "VAF",
    "Location", "Location2", "AFmax", "controlled_genes", "diag_hboc",
    "occurrences_flowcell", "occurrences_pos", "occurrences_strict",
    "promoteur_BRCA", "Recherche", "AnnotSV_ID",
    "SV_chrom", "SV_start", "SV_end", "FILTER", "FORMAT",
    "Annotation_mode", "CytoBand", "Closest_left", "Closest_right",
    "Gene_count", "Tx", "Tx_version", "Tx_start", "Tx_end",
    "Overlapped_tx_length", "Overlapped_CDS_length", "Overlapped_CDS_percent",
    "Frameshift", "Exon_count", "Dist_nearest_SS", "Nearest_SS_type",
    "Intersect_start", "Intersect_end", "RE_gene",
    "P_gain_phen", "P_gain_hpo", "P_gain_source", "P_gain_coord",
    "P_loss_phen", "P_loss_hpo", "P_loss_source", "P_loss_coord",
    "P_ins_phen", "P_ins_hpo", "P_ins_source", "P_ins_coord",
    "po_P_gain_phen", "po_P_gain_hpo", "po_P_gain_source", "po_P_gain_coord",
    "po_P_gain_percent",
    "po_P_loss_phen", "po_P_loss_hpo", "po_P_loss_source", "po_P_loss_coord",
    "po_P_loss_percent",
    "P_snvindel_nb", "P_snvindel_phen",
    "B_gain_source", "B_gain_coord", "B_gain_AFmax",
    "B_loss_source", "B_loss_coord", "B_loss_AFmax",
    "B_ins_source", "B_ins_coord", "B_ins_AFmax",
    "B_inv_source", "B_inv_coord", "B_inv_AFmax",
    "po_B_gain_allG_source", "po_B_gain_allG_coord",
    "po_B_gain_someG_source", "po_B_gain_someG_coord",
    "po_B_loss_allG_source", "po_B_loss_allG_coord",
    "po_B_loss_someG_source", "po_B_loss_someG_coord",
    "GC_content_left", "GC_content_right",
    "Repeat_coord_left", "Repeat_type_left",
    "Repeat_coord_right", "Repeat_type_right",
    "Gap_left", "Gap_right", "SegDup_left", "SegDup_right",
    "ENCODE_blacklist_left", "ENCODE_blacklist_characteristics_left",
    "ENCODE_blacklist_right", "ENCODE_blacklist_characteristics_right",
    "ACMG", "HI", "TS", "DDD_HI_percent",
    "ExAC_delZ", "ExAC_dupZ", "ExAC_cnvZ", "ExAC_synZ",
    "GenCC_disease", "GenCC_moi", "GenCC_classification", "GenCC_pmid",
    "NCBI_gene_ID", "OMIM_ID", "OMIM_phenotype", "OMIM_inheritance",
    "OMIM_morbid", "OMIM_morbid_candidate",
    "LOEUF_bin", "GnomAD_pLI", "ExAC_pLI",
    "AnnotSV_ranking_score", "AnnotSV_ranking_criteria", "ACMG_class",
    "INFO", "SVTYPE", "SVLEN", "END", "STRAND", "STDEV_LEN", "STDEV_POS",
    "greendb_id", "greendb_stdtype", "greendb_dbsource", "greendb_genes",
    "greendb_constraint", "ANN", "GT", "GQ", "DR", "DV",
    "SUPPORT_LONG", "CHR2"
]


def open_text(path):
    """Open plain text or .gz text files."""
    if str(path).endswith(".gz"):
        return gzip.open(path, "rt")
    return open(path, "rt")


def normalize_chr(chrom):
    """Return chromosome without leading 'chr', for robust comparisons."""
    s = str(chrom).strip()
    if s.lower().startswith("chr"):
        s = s[3:]
    return s


def ensure_chr(chrom):
    """Return chromosome with one leading 'chr'."""
    s = str(chrom).strip()
    if not s:
        return s
    if s.lower().startswith("chr"):
        return "chr" + s[3:]
    return "chr" + s


def split_info(info):
    """
    Split the INFO column.

    Fix vs previous version:
    - the first INFO field is not discarded anymore if it is key=value.
    - flag fields like PRECISE / IMPRECISE are kept in INFO.
    """
    new_columns = {}
    flags = []

    if pd.isna(info):
        return new_columns

    for field in str(info).split(";"):
        field = field.strip()
        if not field or field == ".":
            continue

        if "=" in field:
            key, value = field.split("=", 1)
            key = key.strip()
            if key:
                new_columns[key] = value
        else:
            flags.append(field)

    if flags:
        new_columns["INFO"] = ";".join(flags)

    return new_columns


def normalize_sv_read_metrics(df):
    """Expose cuteSV counts in report columns without inventing missing coverage."""
    if "RE" not in df.columns:
        return df

    def fill_missing(column, values):
        if column not in df.columns:
            df[column] = ""
        missing = df[column].isna() | df[column].astype(str).str.lower().isin(["", ".", "nan", "none"])
        # pandas 3 infers a strict string dtype for newly created columns.
        # Use object here because report columns can contain numbers and blanks.
        df[column] = df[column].astype(object)
        df.loc[missing, column] = values.loc[missing].astype(object).fillna("")

    fill_missing("SUPPORT", pd.to_numeric(df["RE"], errors="coerce"))
    if "AF" in df.columns:
        fill_missing("VAF", pd.to_numeric(df["AF"], errors="coerce"))
    if {"DR", "DV"}.issubset(df.columns):
        ref = pd.to_numeric(df["DR"], errors="coerce")
        alt = pd.to_numeric(df["DV"], errors="coerce")
        coverage = (ref + alt).where((ref >= 0) & (alt >= 0))
        fill_missing("COVERAGE", coverage)
        fill_missing("VAF", alt / coverage.where(coverage > 0))
    return df


def read_annotsv_tsv(path):
    """
    Read an AnnotSV-like TSV robustly.

    Handles:
    - optional metadata/comment lines before the header
    - a header column starting with '#'
    - header-only/comment-only files as empty inputs
    - .tsv.gz files
    """
    expected = {"SV_chrom", "SV_start", "SV_end", "SV_type"}
    header_idx = None
    has_data_lines = False

    with open_text(path) as handle:
        for i, line in enumerate(handle):
            stripped = line.rstrip("\n")
            if not stripped:
                continue

            if not stripped.startswith("#"):
                has_data_lines = True

            candidate = stripped.lstrip("#")
            cols = [c.strip() for c in candidate.split("\t")]

            if expected.intersection(cols) or "AnnotSV_ID" in cols or "Gene_name" in cols:
                header_idx = i
                break

    if header_idx is None:
        if not has_data_lines:
            return pd.DataFrame()

        # Fallback to the old behavior, but fail later with a clear missing-column error if wrong.
        df = pd.read_csv(path, sep="\t", dtype=str, on_bad_lines="warn")
    else:
        df = pd.read_csv(path, sep="\t", dtype=str, skiprows=header_idx, header=0, on_bad_lines="warn")

    df.columns = [str(c).lstrip("#").strip() for c in df.columns]
    df = df.fillna("")

    # Drop fully empty rows if any.
    if not df.empty:
        non_empty = df.apply(lambda row: any(str(x).strip() for x in row), axis=1)
        df = df.loc[non_empty].copy()

    return df


def merge_info_columns(df):
    """Expand INFO and merge into df without creating duplicate column names."""
    if "INFO" not in df.columns:
        print("[WARN] No INFO column found; continuing without INFO expansion.")
        return df

    info_expanded = df["INFO"].apply(split_info)
    info_df = pd.DataFrame(info_expanded.tolist(), index=df.index).fillna("")

    base = df.drop(columns=["INFO"])

    for col in info_df.columns:
        if col in base.columns:
            # Fill only empty values in the existing column.
            base_str = base[col].fillna("").astype(str)
            info_str = info_df[col].fillna("").astype(str)
            base[col] = np.where(base_str.str.strip().ne(""), base[col], info_str)
        else:
            base[col] = info_df[col]

    return base


def require_columns(df, input_file, required_cols):
    missing = [c for c in required_cols if c not in df.columns]
    if missing:
        raise ValueError(
            f"Missing required column(s) in {input_file}: {', '.join(missing)}. "
            f"Available columns: {', '.join(map(str, df.columns[:40]))}"
        )


def ensure_optional_columns(df, columns, default=""):
    for col in columns:
        if col not in df.columns:
            df[col] = default
    return df


def contains_target_gene(value, genes=TARGET_GENE_SET):
    """Exact-ish token match on common separators."""
    if pd.isna(value):
        return False
    text = str(value).strip()
    if not text or text in {".", "nan", "None"}:
        return False
    tokens = re.split(r"[;,|\s]+", text)
    return any(t.strip() in genes for t in tokens if t.strip())


def annotate_promoters(df, bed_file):
    """
    Adds 'promoteur_BRCA' based on overlap with promoter regions in a BED file.
    Overlap rule: SV_start ∈ [START, END)

    Fix vs previous version:
    - chromosome comparison is normalized, so '1' and 'chr1' both match.
    """
    df["promoteur_BRCA"] = "no"

    if not os.path.exists(bed_file):
        print(f"[WARN] Promoter BED not found: {bed_file}. promoteur_BRCA set to 'no'.")
        return df

    promoter_df = pd.read_csv(
        bed_file,
        sep="\t",
        header=None,
        names=["CHROM", "START", "END", "ANNOTATION"],
        dtype={"CHROM": str},
    )

    sv_chrom_norm = df["SV_chrom"].apply(normalize_chr)
    sv_start_num = pd.to_numeric(df["SV_start"], errors="coerce")

    for _, row in promoter_df.iterrows():
        chrom = normalize_chr(row["CHROM"])
        start = int(row["START"])
        end = int(row["END"])
        annotation = row["ANNOTATION"]

        mask = (
            (sv_chrom_norm == chrom) &
            (sv_start_num >= start) &
            (sv_start_num < end)
        )
        df.loc[mask, "promoteur_BRCA"] = annotation

    return df


def add_occurrences_columns(df, database_file, sv=False):
    """
    Read a TSV DB with columns: ID, sample_ids comma-separated.

    - sv=True  → strict key: ID_ALT = ID + "_" + ALT
    - sv=False → positional key: base ID
    """
    output_column = "occurrences_strict" if sv else "occurrences_pos"

    if df.empty:
        df[output_column] = pd.Series(dtype="Int64")
        return df

    if not database_file or not os.path.exists(database_file):
        df[output_column] = 0
        return df

    db_df = pd.read_csv(database_file, sep="\t", dtype=str).fillna("")
    if "ID" not in db_df.columns or "sample_ids" not in db_df.columns:
        df[output_column] = 0
        return df

    db_df["ID"] = db_df["ID"].astype(str).str.strip()
    db_df["sample_ids"] = db_df["sample_ids"].apply(
        lambda s: set(t.strip() for t in s.split(",") if t.strip())
        if isinstance(s, str) and s else set()
    )

    # IDs are expected as base or base_ALT, for example:
    # chr1:100-200_DEL or chr1:100-200_DEL_<ALT>
    strict_pat = re.compile(r"_(DEL|DUP|INS|INV|BND|TRA)_[^_]+$")

    strict_map = {}
    base_union = {}

    for k, sids in zip(db_df["ID"], db_df["sample_ids"]):
        if strict_pat.search(k):
            base = k.rsplit("_", 1)[0]
            strict_map[k] = sids
        else:
            base = k
        base_union.setdefault(base, set()).update(sids)

    def minus_current(samples, current):
        if current and current in samples:
            return len(samples) - 1
        return len(samples)

    def count_for_row(row):
        base_id = str(row.get("ID", "")).strip()
        alt = str(row.get("ALT", "")).strip()
        current = str(row.get("Sample_ID", "")).strip()

        if sv:
            key = f"{base_id}_{alt}" if alt else base_id
            samples = strict_map.get(key, base_union.get(base_id, set()))
            return minus_current(samples, current)

        samples = base_union.get(base_id, set())
        return minus_current(samples, current)

    df[output_column] = df.apply(count_for_row, axis=1).astype(int)
    return df


# ---------- Number handling & Excel formatting ----------

def _coerce_numeric_columns(df, min_numeric_ratio=0.9, force_cols=None):
    """
    Try to convert object columns to numeric.
    Accept both '.' and ',' decimal separators.
    """
    force_cols = set(force_cols or [])
    float_cols, int_cols = set(), set()
    out = df.copy()

    for col in out.columns:
        s_orig = out[col]

        if pd.api.types.is_numeric_dtype(s_orig):
            if pd.api.types.is_float_dtype(s_orig):
                float_cols.add(col)
            elif pd.api.types.is_integer_dtype(s_orig):
                int_cols.add(col)
            continue

        if s_orig.dtype == object:
            s_str = s_orig.fillna("").astype(str)
            s_str = s_str.replace({"nan": "", "None": "", ".": ""})

            non_empty_mask = s_str.str.strip().ne("")
            non_empty = int(non_empty_mask.sum())
            if non_empty == 0:
                continue

            s_try = pd.to_numeric(s_str.str.replace(",", ".", regex=False), errors="coerce")
            numeric_count = int(s_try.notna().sum())
            required_ratio = 0.0 if col in force_cols else min_numeric_ratio

            if numeric_count > 0 and (numeric_count / non_empty) >= required_ratio:
                s_non_na = s_try.dropna()
                if len(s_non_na) and (np.mod(s_non_na, 1) == 0).all():
                    out[col] = s_try.round().astype("Int64")
                    int_cols.add(col)
                else:
                    out[col] = s_try.astype(float)
                    float_cols.add(col)

    return out, sorted(float_cols), sorted(int_cols)


def _write_excel_fr_numbers(df, path, sheet_name="Sheet1"):
    """Write DataFrame to Excel with true numeric cells and French-looking decimals."""
    from openpyxl.utils import get_column_letter

    df_to_write = df.copy()
    if df_to_write.empty and len(df_to_write.columns) == 0:
        # Avoid a totally blank workbook if something unexpected happens.
        df_to_write = pd.DataFrame({"message": ["No rows to write"]})

    df_conv, float_cols, int_cols = _coerce_numeric_columns(df_to_write, force_cols={"AFmax"})

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df_conv.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]

        float_fmt = "#,##0.############"
        int_fmt = "#,##0"
        max_row = ws.max_row

        for col_name in float_cols:
            col_idx = df_conv.columns.get_loc(col_name) + 1
            col_letter = get_column_letter(col_idx)
            for r in range(2, max_row + 1):
                ws[f"{col_letter}{r}"].number_format = float_fmt

        for col_name in int_cols:
            col_idx = df_conv.columns.get_loc(col_name) + 1
            col_letter = get_column_letter(col_idx)
            for r in range(2, max_row + 1):
                ws[f"{col_letter}{r}"].number_format = int_fmt

        ws.freeze_panes = "A2"
        last_col_letter = get_column_letter(ws.max_column)
        ws.auto_filter.ref = f"A1:{last_col_letter}{ws.max_row}"


# ========================= Main pipeline =========================

def process_annotsv(input_files, output_filtered, output_unfiltered,
                    notes_file, bed_file, occurrence_file, occurrence_sv_file):
    """
    Process AnnotSV TSV files and write filtered/unfiltered Excel outputs.

    Important fix:
    - This function no longer catches every exception and writes empty Excel files.
      Real errors now fail the job, which prevents silent empty outputs.
    """
    if not input_files:
        raise ValueError("No input TSV files were provided.")

    df_filtered_all = []
    df_unfiltered_all = []

    # Load notes; accept several possible column names but output with new names.
    if os.path.exists(notes_file):
        df_notes = pd.read_csv(notes_file, sep="\t", dtype=str).fillna("")
    else:
        print(f"[WARN] Notes file not found: {notes_file}. Continuing without previous notes.")
        df_notes = pd.DataFrame(columns=["ID"])

    note_candidates = ["notes_variant", "notes", "anciennes_notes", "pré-notes_variant"]
    note_col = next((c for c in note_candidates if c in df_notes.columns), None)

    if "ID" not in df_notes.columns:
        df_notes["ID"] = ""

    df_notes["__note_merge"] = df_notes[note_col] if note_col else ""
    df_notes = df_notes[["ID", "__note_merge"]].drop_duplicates(subset=["ID"], keep="first")

    # Prefer the dedicated SV database, but fall back to occurrence_file if needed.
    sv_db_for_occurrences = occurrence_sv_file if occurrence_sv_file else occurrence_file

    for input_file in input_files:
        print(f"[INFO] Reading {input_file}")
        df = read_annotsv_tsv(input_file)

        if df.empty:
            print(f"[WARN] Skipping file {input_file}: no variant lines found.")
            continue

        print(f"[INFO] Loaded {len(df)} rows and {len(df.columns)} columns from {input_file}")

        # Required for stable SV ID and promoter annotation.
        require_columns(df, input_file, ["SV_chrom", "SV_start", "SV_end", "SV_type"])

        # Derive sample name from filename prefix before first underscore.
        file_prefix = os.path.basename(input_file).strip().split("_")[0]
        df["Sample_ID"] = file_prefix

        # Remove some columns if present.
        columns_to_remove = ["SUPPORT_LONG", "Samples_ID", file_prefix]
        df.drop(columns=[c for c in columns_to_remove if c in df.columns], inplace=True, errors="ignore")

        # Expand INFO robustly if present. Missing INFO is no longer fatal.
        df = normalize_sv_read_metrics(merge_info_columns(df))

        # Gene columns are useful for filtering but should not make the whole job fail.
        missing_gene_cols = [c for c in ["Gene_name", "controlled_genes"] if c not in df.columns]
        if missing_gene_cols:
            print(f"[WARN] {input_file}: missing optional gene column(s): {', '.join(missing_gene_cols)}. Filling with empty values.")
        df = ensure_optional_columns(df, ["Gene_name", "controlled_genes"], default="")

        # Build stable SV ID.
        df["ID"] = (
            df["SV_chrom"].apply(ensure_chr).astype(str) + ":" +
            df["SV_start"].astype(str) + "-" +
            df["SV_end"].astype(str) + "_" +
            df["SV_type"].astype(str)
        )

        # occurrences_flowcell: keep your original ceil(count/2) logic.
        df["occurrences_flowcell"] = np.ceil(df["ID"].map(df["ID"].value_counts()) / 2).astype(int)

        # Merge notes → new column name 'pré-notes_variant'.
        df = df.merge(df_notes, on="ID", how="left")
        df.rename(columns={"__note_merge": "pré-notes_variant"}, inplace=True)
        df["pré-notes_variant"] = df["pré-notes_variant"].fillna("")

        # Manual annotation columns.
        df = ensure_optional_columns(df, ["notes_variant", "notes_échantillon", "Validation", "Recherche"], default="")

        # diag_hboc flag.
        df["diag_hboc"] = df["Gene_name"].apply(
            lambda x: "oui" if contains_target_gene(x, TARGET_GENE_SET) else "non"
        )

        # AFmax by SV_type if present.
        afmax_map = {
            "DEL": "B_loss_AFmax",
            "DUP": "B_gain_AFmax",
            "INS": "B_ins_AFmax",
            "INV": "B_inv_AFmax",
        }

        def get_afmax(row):
            svt = str(row.get("SV_type", "")).strip()
            col = afmax_map.get(svt)
            if col and col in df.columns:
                return row.get(col, "")
            return ""

        df["AFmax"] = df.apply(get_afmax, axis=1)

        # Promoter annotation.
        df = annotate_promoters(df, bed_file)

        # Occurrence counts from SV DB.
        df = add_occurrences_columns(df, sv_db_for_occurrences, sv=False)  # -> occurrences_pos
        df = add_occurrences_columns(df, sv_db_for_occurrences, sv=True)   # -> occurrences_strict

        # Ensure leading columns exist.
        df = ensure_optional_columns(df, ["pré-notes_variant", "notes_variant", "notes_échantillon"], default="")

        # Reorder: requested columns first, then append any remaining.
        ordered = [c for c in DESIRED_ORDER if c in df.columns]
        remaining = [c for c in df.columns if c not in ordered]
        df = df[ordered + remaining]

        # Filtered view: Gene_name or controlled_genes intersects targets.
        gene_mask = df["Gene_name"].apply(lambda x: contains_target_gene(x, TARGET_GENE_SET))
        controlled_mask = df["controlled_genes"].apply(lambda x: contains_target_gene(x, TARGET_GENE_SET))
        df_filtered = df[gene_mask | controlled_mask].copy()

        print(f"[INFO] {input_file}: unfiltered={len(df)}, filtered={len(df_filtered)}")

        df_unfiltered_all.append(df)
        df_filtered_all.append(df_filtered)

    if not df_unfiltered_all:
        raise RuntimeError(
            "No non-empty AnnotSV rows were processed from any input file. "
            "Check that the script receives AnnotSV TSV files, not raw VCF files, and that upstream AnnotSV outputs are not header-only."
        )

    df_unfiltered_final = pd.concat(df_unfiltered_all, ignore_index=True)

    if df_filtered_all:
        df_filtered_final = pd.concat(df_filtered_all, ignore_index=True)
    else:
        df_filtered_final = df_unfiltered_final.iloc[0:0].copy()

    # If filtered is empty, preserve the same columns as unfiltered.
    if df_filtered_final.empty and len(df_filtered_final.columns) == 0:
        df_filtered_final = df_unfiltered_final.iloc[0:0].copy()

    print(f"[INFO] Final unfiltered rows: {len(df_unfiltered_final)}")
    print(f"[INFO] Final filtered rows: {len(df_filtered_final)}")

    _write_excel_fr_numbers(df_filtered_final, output_filtered, sheet_name="Filtered")
    _write_excel_fr_numbers(df_unfiltered_final, output_unfiltered, sheet_name="Unfiltered")

    print("✅ Processed AnnotSV files saved:")
    print(f"  - {output_filtered}")
    print(f"  - {output_unfiltered}")


# ========================= CLI =========================

if __name__ == "__main__":
    # Args:
    #   one or more input .tsv files
    #   output_filtered.xlsx
    #   output_unfiltered.xlsx
    #   notes.tsv
    #   promoters.bed
    #   variants_db.tsv
    #   sv_variants_db.tsv
    #
    # Last 6 args are fixed in order.
    if len(sys.argv) < 8:
        raise SystemExit(
            "Usage: python script.py <input1.tsv> [input2.tsv ...] "
            "<output_filtered.xlsx> <output_unfiltered.xlsx> "
            "<notes.tsv> <promoters.bed> <variants_db.tsv> <sv_variants_db.tsv>"
        )

    input_files = sys.argv[1:-6]
    output_filtered = sys.argv[-6]
    output_unfiltered = sys.argv[-5]
    notes_file = sys.argv[-4]
    bed_file = sys.argv[-3]
    occurrence_file = sys.argv[-2]
    occurrence_sv_file = sys.argv[-1]

    try:
        process_annotsv(
            input_files,
            output_filtered,
            output_unfiltered,
            notes_file,
            bed_file,
            occurrence_file,
            occurrence_sv_file,
        )
    except Exception:
        print("\nERROR: AnnotSV processing failed. No fake empty Excel outputs were written.", file=sys.stderr)
        traceback.print_exc()
        sys.exit(1)
