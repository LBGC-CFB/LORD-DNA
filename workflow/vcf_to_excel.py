#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import sys
import pandas as pd
import numpy as np
import re
import os

# ------------------------------ Constants ------------------------------

# SpliceAI prediction columns
pred_cols = ['SpliceAI_pred_DS_AG', 'SpliceAI_pred_DS_AL', 'SpliceAI_pred_DS_DG', 'SpliceAI_pred_DS_DL']
ds_cols   = ['DS_AG', 'DS_AL', 'DS_DG', 'DS_DL']

# Gene lists for filtering
list_gene = ['BRCA1', 'BRCA2', 'RAD51C', 'RAD51D', 'PMS2', 'PALB2', 'PTEN', 'TP53', 'CDH1', 'MLH1', 'MSH2', 'MSH6']
DRGFP = ['BRCA1', 'BRCA2', 'RAD51C', 'RAD51D', 'PALB2']

# Columns to write as *true numbers* in Excel, with French-looking decimals (via cell formats)
columns_to_convert = (
    ['AF','QUAL','am_pathogenicity','SpliceAI_MAX']
    + pred_cols + ds_cols
    + [
        'gnomADg_AF', 'gnomADg_AFR_AF','gnomADg_AMI_AF', 'gnomADg_AMR_AF', 'gnomADg_ASJ_AF',
        'gnomADg_EAS_AF','gnomADg_FIN_AF', 'gnomADg_MID_AF', 'gnomADg_NFE_AF', 'gnomADg_OTH_AF',
        'gnomADg_SAS_AF','greendb_constraint', 'MaxEntScan_ref',
    ]
)

# Desired column ordering
essential_cols = [
    'pré-notes_variant', 'notes_variant', 'notes_échantillon',
    'Validation','Sample_ID','SYMBOL','Feature','EXON','INTRON','HGVSc','HGVSp','HGVSg',
    'REF','ALT','EventSize','AF','DP','AD','FILTER','QUAL','IMPACT','Consequence',
    'CLIN_SIG','Existing_variation','gnomADg_AF','MANE_SELECT','SpliceAI_MAX','SPiP_Interpretation',
    'SPiP_InterConfident','SIFT','am_class','am_pathogenicity','PolyPhen','REVEL','controlled_genes',
    'greendb_id','greendb_stdtype','greendb_dbsource','greendb_genes','greendb_constraint','greendb_level',
    'Flowcell_Occurrences','Recherche','diag_hboc','diag_filters','DRGFP','promoteur_BRCA','Occurrences_Totales'
]

additional_cols = [
    'gnomADg_AFR_AF','gnomADg_AMI_AF','gnomADg_AMR_AF','gnomADg_ASJ_AF','gnomADg_EAS_AF',
    'gnomADg_FIN_AF','gnomADg_MID_AF','gnomADg_NFE_AF','gnomADg_OTH_AF','gnomADg_SAS_AF',
    'SpliceAI_pred_DP_AG','SpliceAI_pred_DP_AL','SpliceAI_pred_DP_DG','SpliceAI_pred_DP_DL',
    'SpliceAI_pred_DS_AG','SpliceAI_pred_DS_AL','SpliceAI_pred_DS_DG','SpliceAI_pred_DS_DL',
    'SpliceAI_pred_SYMBOL','NAME','DS_AG','DS_AL','DS_DG','DS_DL','DP_AG','DP_AL','DP_DG','DP_DL',
    'DS_AG_REF','DS_AG_ALT','DS_AL_REF','DS_AL_ALT','DS_DG_REF','DS_DG_ALT','DS_DL_REF','DS_DL_ALT',
    'MaxEntScan_diff','MaxEntScan_ref','BIOTYPE','cDNA_position','CDS_position','Gene','DISTANCE',
    'STRAND','ID','GT','GQ'
]

# ------------------------------ Helpers ------------------------------

def get_spliceai_max(row):
    """Return max SpliceAI among pred_cols or ds_cols (numeric)."""
    pred_vals = pd.to_numeric(row.reindex(pred_cols), errors='coerce')
    ds_vals = pd.to_numeric(row.reindex(ds_cols), errors='coerce')
    if pred_vals.notna().any():
        return pred_vals.max()
    if ds_vals.notna().any():
        return ds_vals.max()
    return np.nan

def annotate_promoters(df, bed_file):
    """Adds 'promoteur_BRCA' based on overlap with promoter regions (POS in [START, END))."""
    promoter_df = pd.read_csv(
        bed_file, sep='\t', header=None,
        names=["CHROM", "START", "END", "ANNOTATION"],
        dtype={'CHROM': str}
    )
    df["promoteur_BRCA"] = "no"
    pos_num = pd.to_numeric(df["POS"], errors="coerce")
    for _, r in promoter_df.iterrows():
        mask = (
            (df["CHROM"] == r["CHROM"]) &
            (pos_num >= int(r["START"])) &
            (pos_num <  int(r["END"]))
        )
        df.loc[mask, "promoteur_BRCA"] = r["ANNOTATION"]
    return df

def add_occurrences_totales_column(df, database_file):
    """Adds 'Occurrences_Totales' = count of other samples sharing same ID."""
    if not os.path.exists(database_file):
        df["Occurrences_Totales"] = 0
        return df
    db_df = pd.read_csv(database_file, sep="\t", dtype=str).fillna("")
    if "ID" not in db_df.columns or "sample_ids" not in db_df.columns:
        df["Occurrences_Totales"] = 0
        return df
    db_df["sample_ids"] = db_df["sample_ids"].apply(lambda x: set(x.split(",")) if x else set())
    db_dict = dict(zip(db_df["ID"], db_df["sample_ids"]))
    def count_other_samples(row):
        samples = db_dict.get(row["ID"], set())
        return len(samples - {row["Sample_ID"]})
    df["Occurrences_Totales"] = df.apply(count_other_samples, axis=1)
    return df

def _series_or_nan(df: pd.DataFrame, col: str) -> pd.Series:
    """Return df[col] if exists else a NaN series of same length."""
    if col in df.columns:
        return df[col]
    return pd.Series([np.nan] * len(df), index=df.index)


def _is_missing_value(value) -> bool:
    if value is None:
        return True
    value = str(value).strip()
    return value == "" or value == "." or value.lower() in {"nan", "none", "na"}


def _format_af(value: float) -> str:
    text = f"{value:.6f}".rstrip("0").rstrip(".")
    return text if text else "0"


def compute_af_from_sample(format_keys, sample_values):
    """
    Recover AF from sample FORMAT when INFO/AF is absent.

    Clair3 uses FORMAT/AF, DeepVariant uses FORMAT/VAF. If neither exists,
    calculate ALT depth / DP from AD and DP.
    """
    fmt = {k: v for k, v in zip(format_keys, sample_values)}

    for key in ("AF", "VAF"):
        value = fmt.get(key)
        if not _is_missing_value(value):
            return value

    ad_value = fmt.get("AD")
    if _is_missing_value(ad_value):
        return np.nan

    try:
        depths = [float(x) for x in str(ad_value).split(",") if not _is_missing_value(x)]
    except ValueError:
        return np.nan

    if len(depths) < 2:
        return np.nan

    dp_value = fmt.get("DP")
    try:
        dp = float(dp_value) if not _is_missing_value(dp_value) else sum(depths)
    except ValueError:
        dp = sum(depths)

    if dp <= 0:
        return np.nan

    return ",".join(_format_af(alt_depth / dp) for alt_depth in depths[1:])

# ---- PRE-TAG DIAG FILTERS (adds diag_filters yes/no on filtered output) ----

def build_diag_filter_mask(df: pd.DataFrame) -> pd.Series:
    """Return boolean mask of rows that pass all diagnostic filters (pre-tag logic)."""
    # 1) diag_hboc == "oui"
    diag = (
        _series_or_nan(df, "diag_hboc")
        .astype(str).str.strip().str.lower()
        == "oui"
    )

    # 2) gnomADg_AF <= 0.001 OR empty
    gnom_col = pd.to_numeric(_series_or_nan(df, "gnomADg_AF"), errors="coerce")
    gnom = (gnom_col <= 0.001) | gnom_col.isna()

    # 3) Occurrences_Totales <= 2 (NaN -> False)
    occ_col = pd.to_numeric(_series_or_nan(df, "Occurrences_Totales"), errors="coerce")
    occ = occ_col <= 2

    # 4) big OR section
    spliceai_col = pd.to_numeric(_series_or_nan(df, "SpliceAI_MAX"), errors="coerce")
    cond_spliceai = spliceai_col >= 0.1

    spip_col = _series_or_nan(df, "SPiP_Interpretation").astype(str).str.strip()
    cond_spip = spip_col != "NTR"  # NaN -> "nan" -> kept

    prom_col = _series_or_nan(df, "promoteur_BRCA").astype(str).str.strip().str.lower()
    cond_prom = prom_col != "no"

    eventsize_col = pd.to_numeric(_series_or_nan(df, "EventSize"), errors="coerce")
    cond_eventsize = eventsize_col > 10

    impact_col = _series_or_nan(df, "IMPACT").astype(str).str.strip().str.upper()
    cond_impact = impact_col == "HIGH"

    cons_col = _series_or_nan(df, "Consequence").astype(str).str.lower()
    cond_cons = cons_col.str.contains(
        "frameshift|nonsense|missense|splice_region",
        regex=True,
        na=False,
    )

    big_or = cond_spliceai | cond_spip | cond_prom | cond_eventsize | cond_impact | cond_cons

    return diag & gnom & occ & big_or

def _move_col_after(df: pd.DataFrame, col: str, after: str) -> pd.DataFrame:
    """Move a column right after another column if both exist."""
    if col not in df.columns or after not in df.columns:
        return df
    cols = list(df.columns)
    cols.remove(col)
    idx = cols.index(after) + 1
    cols.insert(idx, col)
    return df[cols]

# ---- Excel writing with *true numbers* + French-looking decimals ----

def _coerce_specific_numeric(df, cols):
    """Coerce only selected columns to numeric, accepting ',' or '.' decimals."""
    out = df.copy()
    float_cols, int_cols = [], []
    for col in cols:
        if col not in out.columns:
            continue
        s = out[col]
        if pd.api.types.is_numeric_dtype(s):
            if pd.api.types.is_float_dtype(s):
                float_cols.append(col)
            elif pd.api.types.is_integer_dtype(s):
                int_cols.append(col)
            continue
        s_str = s.astype(str)
        s_num = pd.to_numeric(s_str.str.replace(",", ".", regex=False), errors="coerce")
        s_non_na = s_num.dropna()
        if len(s_non_na) and np.allclose(s_non_na, np.round(s_non_na)):
            out[col] = s_num.round().astype("Int64")
            int_cols.append(col)
        else:
            out[col] = s_num.astype(float)
            float_cols.append(col)
    return out, float_cols, int_cols

def write_excel_fr(path, sheet_name, df, float_cols, int_cols):
    """Write df to Excel with number formats so Excel shows comma decimals in FR locales (openpyxl-only)."""
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]

        fltfmt = "#,##0.############"
        intfmt = "#,##0"

        max_row = ws.max_row

        # Apply formats to float columns
        for c in float_cols:
            if c in df.columns:
                j = df.columns.get_loc(c) + 1
                col_letter = get_column_letter(j)
                for r in range(2, max_row + 1):
                    ws[f"{col_letter}{r}"].number_format = fltfmt

        # Apply formats to int columns
        for c in int_cols:
            if c in df.columns:
                j = df.columns.get_loc(c) + 1
                col_letter = get_column_letter(j)
                for r in range(2, max_row + 1):
                    ws[f"{col_letter}{r}"].number_format = intfmt

        # Freeze header & add autofilter
        ws.freeze_panes = "A2"
        last_col_letter = get_column_letter(ws.max_column)
        ws.auto_filter.ref = f"A1:{last_col_letter}{ws.max_row}"

# ------------------------------ Main ------------------------------

def process_vcf_files(vcf_files, output_filtered, output_unfiltered, notes_file, bed_file, occurrence_file):
    """Parse VCFs, annotate, and export Excel with correct typing/formatting."""
    vcf_data = []
    info_fields = {}
    split_subfields, splice_subfields = [], []

    # Parse headers to capture INFO subfield formats
    for vcf_input in vcf_files:
        file_prefix = os.path.basename(vcf_input).split('.')[0].split('_')[0].strip()
        #file_prefix = "_".join(os.path.basename(vcf_input).split('.')[0].split('_')[0:2]).strip()
        ann_columns, splice_columns = [], []
        with open(vcf_input) as infile:
            for line in infile:
                if line.startswith("##INFO"):
                    m = re.search(r'##INFO=<ID=([^,]+),.*Description="([^"]+)', line)
                    if m:
                        field_id, _ = m.groups()
                        info_fields[field_id] = field_id
                        if field_id == "ANN":
                            if "Format: " in line:
                                ann_columns = line.split("Format: ")[-1].strip().split("|")
                        elif field_id == "SpliceAI":
                            if "Format: " in line:
                                splice_columns = line.split("Format: ")[-1].strip().split("|")
                elif line.startswith("#CHROM"):
                    column_names = line.strip().lstrip("#").split("\t")
                    if "INFO" in column_names:
                        column_names.remove("INFO")
                    break

        # Parse records
        with open(vcf_input) as infile:
            for line in infile:
                if line.startswith("#"):
                    continue
                fields = line.strip().split("\t")
                base_data = dict(zip(column_names, fields[:7]))
                base_data['Sample_ID'] = f"{file_prefix}"

                info_data, ann_data, splice_data = {}, {}, {}
                for entry in fields[7].split(";"):
                    kv = entry.split("=")
                    if len(kv) != 2:
                        continue
                    key, value = kv
                    if key == "ANN" and ann_columns:
                        for i, subval in enumerate(value.split("|")):
                            if i < len(ann_columns) and ann_columns[i] != "AF":
                                ann_data[ann_columns[i]] = subval
                                if ann_columns[i] not in split_subfields:
                                    split_subfields.append(ann_columns[i])
                    elif key == "SpliceAI" and splice_columns:
                        parts = value.split(",")[0].split("|")  # first transcript
                        for j, subval in enumerate(parts):
                            if j < len(splice_columns):
                                splice_data[splice_columns[j]] = subval
                                if splice_columns[j] not in splice_subfields:
                                    splice_subfields.append(splice_columns[j])
                    elif key in info_fields:
                        info_data[key] = value

                base_data.update(info_data)
                base_data.update(ann_data)
                base_data.update(splice_data)

                if "AF" not in base_data and len(fields) >= 10:
                    format_keys = fields[8].split(":")
                    sample_values = fields[9].split(":")
                    af_value = compute_af_from_sample(format_keys, sample_values)
                    if not pd.isna(af_value):
                        base_data["AF"] = af_value

                vcf_data.append(base_data)

    # Construct DataFrame
    df = pd.DataFrame(vcf_data)
    for col in split_subfields + splice_subfields:
        if col not in df.columns:
            df[col] = np.nan

    # Additional annotations
    df["SpliceAI_MAX"] = df.apply(get_spliceai_max, axis=1)
    df = df.rename(columns={"Interpretation": "SPiP_Interpretation", "InterConfident": "SPiP_InterConfident"})
    df["EventSize"] = df.apply(
        lambda row: max(abs(len(str(row["REF"])) - len(alt)) + 1 for alt in str(row["ALT"]).split(",")),
        axis=1
    )
    df["ID"] = df["CHROM"].astype(str) + ":" + df["POS"].astype(str) + "_" + df["REF"].astype(str) + "_" + df["ALT"].astype(str)
    df["Flowcell_Occurrences"] = df["ID"].map(df["ID"].value_counts()).astype("Int64")

    # Merge notes TSV → output as 'pré-notes_variant'
    df_notes = pd.read_csv(notes_file, sep='\t', dtype=str).fillna("")
    note_candidates = ["notes_variant", "notes", "anciennes_notes", "pré-notes_variant"]
    note_col = next((c for c in note_candidates if c in df_notes.columns), None)
    if "ID" not in df_notes.columns:
        df_notes["ID"] = ""
    df_notes["__note_merge"] = df_notes[note_col] if note_col else ""
    df = df.merge(df_notes[["ID","__note_merge"]], on="ID", how="left").rename(columns={"__note_merge": "pré-notes_variant"})
    df["pré-notes_variant"] = df["pré-notes_variant"].fillna("")

    # Manual annotation columns
    for c in ("notes_variant", "notes_échantillon", "Validation", "Recherche"):
        if c not in df.columns:
            df[c] = ""

    # Clean / compute other fields
    genes = set(list_gene)
    if "SYMBOL" in df.columns:
        df["diag_hboc"] = (
            df["SYMBOL"].fillna("")
            .str.split(r"[;,|]")
            .apply(lambda toks: "oui" if any(t.strip() in genes for t in toks) else "non")
        )
    else:
        df["diag_hboc"] = "non"

    if "controlled_genes" in df.columns:
        df["DRGFP"] = df["controlled_genes"].fillna("").apply(lambda g: "oui" if any(x in g for x in DRGFP) else "non")
    else:
        df["DRGFP"] = "non"

    if "AD" in df.columns:
        df["AD"] = df["AD"].astype(str).str.replace(",", "|", regex=False)
    if "HGVSp" in df.columns:
        df["HGVSp"] = df["HGVSp"].astype(str).apply(lambda x: x.split(":")[-1] if ":" in x else x)

    # Promoters & occurrences
    df = annotate_promoters(df, bed_file)
    df = add_occurrences_totales_column(df, occurrence_file)

    # Final column selection and order (keep any existing subset; append remaining at end)
    final_columns = essential_cols + additional_cols
    present = [c for c in final_columns if c in df.columns]
    remaining = [c for c in df.columns if c not in present]
    df = df[present + remaining]

    # Filtered subset (genes of interest)
    df_filtered = df[
        df['SYMBOL'].fillna('').apply(lambda x: any(g in x for g in list_gene)) |
        df.get('controlled_genes', pd.Series([""]*len(df))).fillna('')
          .apply(lambda x: any(g in x.split(',') for g in list_gene))
    ].copy()

    # ------------------ ADD diag_filters (PRE-TAG) ONLY ON FILTERED ------------------
    mask_diag = build_diag_filter_mask(df_filtered)
    df_filtered["diag_filters"] = np.where(mask_diag, "yes", "no")
    df_filtered = _move_col_after(df_filtered, "diag_filters", "diag_hboc")

    # -------- Excel export: keep true numbers, French-looking decimals --------
    df_u_conv, u_floats, u_ints = _coerce_specific_numeric(df, columns_to_convert)
    df_f_conv, f_floats, f_ints = _coerce_specific_numeric(df_filtered, columns_to_convert)

    write_excel_fr(output_unfiltered, "Unfiltered", df_u_conv, u_floats, u_ints)
    write_excel_fr(output_filtered,  "Filtered",   df_f_conv, f_floats, f_ints)

    print(f"✅ Processed VCFs saved:\n  - {output_unfiltered} (unfiltered)\n  - {output_filtered} (filtered, with diag_filters yes/no)")

# ------------------------------ CLI ------------------------------

if __name__ == "__main__":
    # Usage: python script.py <vcf1> [vcf2 ...] <output_filtered.xlsx> <output_unfiltered.xlsx> <notes.tsv> <promoters.bed> <variants_db.tsv>
    if len(sys.argv) < 7:
        raise SystemExit(
            "Usage: python script.py <vcf1> [vcf2 ...] "
            "<output_filtered.xlsx> <output_unfiltered.xlsx> "
            "<notes.tsv> <promoters.bed> <variants_db.tsv>"
        )
    vcf_files = sys.argv[1:-5]
    output_filtered = sys.argv[-5]
    output_unfiltered = sys.argv[-4]
    notes_file = sys.argv[-3]
    bed_file = sys.argv[-2]
    occurrence_file = sys.argv[-1]
    process_vcf_files(vcf_files, output_filtered, output_unfiltered, notes_file, bed_file, occurrence_file)
